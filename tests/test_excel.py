import io
import uuid
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.db.session import SessionLocal
from app.db.models import RegistroConsumoDB, UsuarioDB

from tests.usuarios import crear_usuario

client = TestClient(app)


def _plan(nombre, kcal, prot, carb, grasa):
    return {"nombre": nombre, "calorias_objetivo": kcal, "proteinas_objetivo": prot,
            "carbohidratos_objetivo": carb, "grasas_objetivo": grasa}


def _usuario_con_consumos():
    """Usuario con plan principal 2300 kcal, un plan de descanso puesto a mano el 2001-05-08 y
    comidas el 7 y el 8 de mayo de 2001 (fechas fijas para no chocar con otros tests)."""
    nombre = f"excel_{uuid.uuid4().hex[:8]}"
    cab = {"Authorization": f"Bearer {crear_usuario(nombre)}"}
    principal = client.get("/api/v1/planes", headers=cab).json()["planes"][0]
    client.put(f"/api/v1/planes/{principal['id']}", headers=cab, json=_plan("Base", 2300, 160, 250, 70))
    descanso = client.post("/api/v1/planes", headers=cab, json=_plan("Descanso", 2000, 150, 200, 65)).json()
    client.put("/api/v1/planes/calendario/2001-05-08", headers=cab, json={"plan_id": descanso["id"]})
    db = SessionLocal()
    try:
        uid = db.query(UsuarioDB).filter(UsuarioDB.username == nombre).first().id
        for fecha, kcal, prot, carb, grasa in [("2001-05-07", 1500, 90, 150, 40), ("2001-05-07", 600, 40.4, 60, 20),
                                               ("2001-05-08", 1990, 149.6, 210, 64)]:
            db.add(RegistroConsumoDB(usuario_id=uid, nombre="Comida", fecha=fecha, calorias=kcal,
                                     proteinas=prot, carbohidratos=carb, grasas=grasa))
        db.commit()
    finally:
        db.close()
    return cab


def test_exportar_excel_por_rango_consumido_sobre_meta():
    cab = _usuario_con_consumos()
    res = client.get("/api/v1/exportar/excel?desde=2001-05-07&hasta=2001-05-09", headers=cab)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "2001-05-07_2001-05-09" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Resumen"]
    filas = [list(f) for f in wb["Resumen"].iter_rows(min_row=4, values_only=True)]
    assert filas[0] == ["Fecha", "Plan", "Calorías (kcal)", "Proteínas (g)", "Carbohidratos (g)", "Grasas (g)"]
    assert len(filas) == 4  # cabecera + una fila por día
    assert filas[1] == ["lun 07/05/2001", "Base", "2100/2300", "130/160", "210/250", "60/70"]
    assert filas[2] == ["mar 08/05/2001", "Descanso", "1990/2000", "150/150", "210/200", "64/65"]
    assert filas[3] == ["mié 09/05/2001", "Base", "0/2300", "0/160", "0/250", "0/70"]


def test_exportar_excel_en_ingles_y_dia_libre():
    cab = _usuario_con_consumos()
    client.put("/api/v1/dias-libres/2001-05-08", headers=cab, json={"libre": True})
    res = client.get("/api/v1/exportar/excel?desde=2001-05-08&hasta=2001-05-08&idioma=en", headers=cab)
    ws = load_workbook(io.BytesIO(res.content))["Summary"]
    assert ws["A4"].value == "Date" and ws["C4"].value == "Calories (kcal)"
    assert ws["A5"].value == "Tue 08/05/2001"
    assert ws["B5"].value == "Descanso · Free day"
    assert ws["C5"].value == "1990/2000"


def test_exportar_excel_rango_invalido():
    assert client.get("/api/v1/exportar/excel?desde=2026-05-10&hasta=2026-05-01").status_code == 400
    assert client.get("/api/v1/exportar/excel?desde=2024-01-01&hasta=2026-01-01").status_code == 400
    assert client.get("/api/v1/exportar/excel?desde=no-es-fecha").status_code == 422
