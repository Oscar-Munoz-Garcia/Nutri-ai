"""Importador masivo de alimentos (despensa) y recetas desde Excel o CSV."""
import csv
import io
import json
import unicodedata
import uuid
from datetime import date
from typing import Dict, List, Optional

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.db.models import PantryDB, RecetaDB

MAX_FILAS = 2000

# Alias aceptados en la cabecera (normalizados: minúsculas, sin acentos, "_" en lugar de espacios)
ALIAS_ALIMENTOS = {
    "nombre": ["nombre", "alimento", "producto", "name"],
    "cantidad_gramos": ["cantidad_gramos", "cantidad", "gramos", "cantidad_g", "peso_g", "g"],
    "calorias_100g": ["calorias_100g", "kcal_100g", "calorias", "kcal", "energia"],
    "proteinas_100g": ["proteinas_100g", "proteinas", "proteina", "prot"],
    "carbohidratos_100g": ["carbohidratos_100g", "carbohidratos", "hidratos", "carbs"],
    "grasas_100g": ["grasas_100g", "grasas", "grasa", "lipidos"],
}

ALIAS_RECETAS = {
    "nombre": ["nombre", "receta", "name"],
    "descripcion": ["descripcion", "instrucciones", "pasos"],
    "calorias_totales": ["calorias_totales", "calorias", "kcal"],
    "proteinas_totales": ["proteinas_totales", "proteinas", "proteina", "prot"],
    "carbohidratos_totales": ["carbohidratos_totales", "carbohidratos", "hidratos", "carbs"],
    "grasas_totales": ["grasas_totales", "grasas", "grasa"],
    "tiempo_preparacion_min": ["tiempo_preparacion_min", "tiempo", "minutos", "tiempo_min"],
    "utensilios_necesarios": ["utensilios_necesarios", "utensilios"],
    "ingredientes": ["ingredientes"],
}

# Columnas cuya presencia indica que el archivo es de recetas
PISTAS_RECETAS = {"calorias_totales", "proteinas_totales", "tiempo_preparacion_min", "tiempo",
                  "utensilios", "utensilios_necesarios", "receta", "ingredientes"}


class ErrorImportacion(ValueError):
    pass


def _normalizar(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    texto = texto.strip().lower()
    for sep in [" ", "-", "/", "(", ")", "."]:
        texto = texto.replace(sep, "_")
    while "__" in texto:
        texto = texto.replace("__", "_")
    return texto.strip("_")


def _numero(valor, por_defecto: float = 0.0) -> float:
    if valor is None or str(valor).strip() == "":
        return por_defecto
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", ".")
    return float(texto)


def _leer_filas(contenido: bytes, nombre_archivo: str) -> List[List]:
    nombre = (nombre_archivo or "").lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise ErrorImportacion("No se pudo leer el archivo Excel. ¿Está dañado o protegido?")
        ws = wb.worksheets[0]
        filas = [list(f) for f in ws.iter_rows(values_only=True)]
        wb.close()
        return filas
    if nombre.endswith((".csv", ".txt", ".tsv")):
        for codificacion in ("utf-8-sig", "latin-1"):
            try:
                texto = contenido.decode(codificacion)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        return [fila for fila in csv.reader(io.StringIO(texto), dialecto)]
    raise ErrorImportacion("Formato no soportado. Usa .xlsx o .csv.")


def _mapear_cabecera(cabecera: List, alias: Dict[str, List[str]]) -> Dict[str, int]:
    normalizada = [_normalizar(c) for c in cabecera]
    mapa = {}
    for campo, opciones in alias.items():
        for opcion in opciones:
            if opcion in normalizada:
                mapa[campo] = normalizada.index(opcion)
                break
    return mapa


def detectar_tipo(cabecera: List) -> str:
    normalizada = {_normalizar(c) for c in cabecera}
    return "recetas" if normalizada & PISTAS_RECETAS else "alimentos"


def _valor(fila: List, mapa: Dict[str, int], campo: str):
    idx = mapa.get(campo)
    if idx is None or idx >= len(fila):
        return None
    return fila[idx]


def _importar_alimentos(db, filas, mapa, usuario_id: int, resultado: Dict) -> None:
    hoy = date.today().isoformat()
    existentes = {
        (p.nombre or "").strip().lower(): p
        for p in db.query(PantryDB).filter(PantryDB.usuario_id == usuario_id).all()
    }
    for n_fila, fila in filas:
        nombre = str(_valor(fila, mapa, "nombre") or "").strip()
        if not nombre:
            resultado["errores"].append({"fila": n_fila, "motivo": "Falta el nombre del alimento."})
            continue
        try:
            datos = {
                "cantidad_gramos": _numero(_valor(fila, mapa, "cantidad_gramos"), 100.0),
                "calorias_100g": _numero(_valor(fila, mapa, "calorias_100g")),
                "proteinas_100g": _numero(_valor(fila, mapa, "proteinas_100g")),
                "carbohidratos_100g": _numero(_valor(fila, mapa, "carbohidratos_100g")),
                "grasas_100g": _numero(_valor(fila, mapa, "grasas_100g")),
            }
        except ValueError:
            resultado["errores"].append({"fila": n_fila, "motivo": f"Valor numérico no válido en '{nombre}'."})
            continue

        item = existentes.get(nombre.lower())
        if item:
            # Mismo criterio que la despensa: se suma la cantidad y se actualizan los valores por 100g
            item.cantidad_gramos = (item.cantidad_gramos or 0) + datos.pop("cantidad_gramos")
            for campo, valor in datos.items():
                setattr(item, campo, valor)
            resultado["actualizados"] += 1
        else:
            item = PantryDB(id=str(uuid.uuid4())[:8], usuario_id=usuario_id, nombre=nombre,
                            fecha_registro=hoy, **datos)
            db.add(item)
            existentes[nombre.lower()] = item
            resultado["importados"] += 1


def _ingredientes_json(valor) -> str:
    """Convierte 'Avena:60; Plátano:100' en JSON [{nombre, cantidad_gramos}]."""
    ingredientes = []
    for parte in str(valor or "").split(";"):
        if not parte.strip():
            continue
        nombre, _, cantidad = parte.partition(":")
        try:
            gramos = _numero(cantidad.lower().replace("g", ""), 0.0)
        except ValueError:
            gramos = 0.0
        ingredientes.append({"nombre": nombre.strip(), "cantidad_gramos": gramos})
    return json.dumps(ingredientes, ensure_ascii=False)


def _importar_recetas(db, filas, mapa, usuario_id: int, resultado: Dict) -> None:
    existentes = {
        (r.nombre or "").strip().lower(): r
        for r in db.query(RecetaDB).filter(RecetaDB.usuario_id == usuario_id).all()
    }
    for n_fila, fila in filas:
        nombre = str(_valor(fila, mapa, "nombre") or "").strip()
        if not nombre:
            resultado["errores"].append({"fila": n_fila, "motivo": "Falta el nombre de la receta."})
            continue
        try:
            datos = {
                "descripcion": str(_valor(fila, mapa, "descripcion") or ""),
                "calorias_totales": _numero(_valor(fila, mapa, "calorias_totales")),
                "proteinas_totales": _numero(_valor(fila, mapa, "proteinas_totales")),
                "carbohidratos_totales": _numero(_valor(fila, mapa, "carbohidratos_totales")),
                "grasas_totales": _numero(_valor(fila, mapa, "grasas_totales")),
                "tiempo_preparacion_min": int(_numero(_valor(fila, mapa, "tiempo_preparacion_min"), 15)),
            }
        except ValueError:
            resultado["errores"].append({"fila": n_fila, "motivo": f"Valor numérico no válido en '{nombre}'."})
            continue
        utensilios = str(_valor(fila, mapa, "utensilios_necesarios") or "Sartén")
        datos["utensilios_necesarios"] = json.dumps(
            [u.strip() for u in utensilios.replace(";", ",").split(",") if u.strip()], ensure_ascii=False)
        datos["ingredientes"] = _ingredientes_json(_valor(fila, mapa, "ingredientes"))

        receta = existentes.get(nombre.lower())
        if receta:
            for campo, valor in datos.items():
                setattr(receta, campo, valor)
            resultado["actualizados"] += 1
        else:
            receta = RecetaDB(id="rec_" + str(uuid.uuid4())[:8], usuario_id=usuario_id, nombre=nombre, **datos)
            db.add(receta)
            existentes[nombre.lower()] = receta
            resultado["importados"] += 1


def importar_archivo(contenido: bytes, nombre_archivo: str, usuario_id: int, tipo: Optional[str] = None) -> Dict:
    """Importa alimentos o recetas. `tipo` = 'alimentos' | 'recetas' | None (autodetectar)."""
    filas = [f for f in _leer_filas(contenido, nombre_archivo) if any(str(c or "").strip() for c in f)]
    if len(filas) < 2:
        raise ErrorImportacion("El archivo no tiene filas de datos (la primera fila debe ser la cabecera).")
    if len(filas) - 1 > MAX_FILAS:
        raise ErrorImportacion(f"Máximo {MAX_FILAS} filas por importación.")

    cabecera, datos = filas[0], filas[1:]
    tipo = tipo or detectar_tipo(cabecera)
    if tipo not in ("alimentos", "recetas"):
        raise ErrorImportacion("Tipo no válido: usa 'alimentos' o 'recetas'.")

    mapa = _mapear_cabecera(cabecera, ALIAS_RECETAS if tipo == "recetas" else ALIAS_ALIMENTOS)
    if "nombre" not in mapa:
        raise ErrorImportacion("No se encontró la columna 'nombre' en la cabecera.")

    resultado = {"tipo": tipo, "importados": 0, "actualizados": 0, "errores": []}
    filas_numeradas = list(enumerate(datos, start=2))
    db = SessionLocal()
    try:
        if tipo == "recetas":
            _importar_recetas(db, filas_numeradas, mapa, usuario_id, resultado)
        else:
            _importar_alimentos(db, filas_numeradas, mapa, usuario_id, resultado)
        db.commit()
    finally:
        db.close()
    return resultado
