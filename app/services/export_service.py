"""Exportador a Excel (.xlsx): una fila por día del rango elegido con lo consumido frente a la meta
de su plan de nutrición, en formato "consumido/meta" (p. ej. 2100/2300)."""
from datetime import date, timedelta
from io import BytesIO
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.db.session import SessionLocal
from app.db.models import RegistroConsumoDB
from app.services import dia_libre_service, plan_service

MAX_DIAS = 366
MACROS = ("calorias", "proteinas", "carbohidratos", "grasas")

COLOR_CABECERA = "4981D4"
COLOR_TITULO = "5749D4"
# Semáforo por celda: verde dentro del ±10% de la meta, ámbar dentro del ±20%, rojo fuera
FILL_OK = PatternFill("solid", start_color="C8EFD4", end_color="C8EFD4")
FILL_AVISO = PatternFill("solid", start_color="FCE8B2", end_color="FCE8B2")
FILL_MAL = PatternFill("solid", start_color="F6C9C9", end_color="F6C9C9")
FILL_LIBRE = PatternFill("solid", start_color="E4E4E7", end_color="E4E4E7")

TEXTOS = {
    "es": {
        "hoja": "Resumen",
        "titulo": "NutriAI — Consumido / meta por día",
        "rango": "Del {desde} al {hasta}",
        "cabecera": ["Fecha", "Plan", "Calorías (kcal)", "Proteínas (g)", "Carbohidratos (g)", "Grasas (g)"],
        "libre": "Día libre",
        "principal": plan_service.NOMBRE_PLAN_PRINCIPAL,
        "dias": ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"],
        "orden": "La fecha de inicio no puede ser posterior a la de fin.",
        "largo": f"El rango no puede superar {MAX_DIAS} días.",
    },
    "en": {
        "hoja": "Summary",
        "titulo": "NutriAI — Consumed / goal per day",
        "rango": "From {desde} to {hasta}",
        "cabecera": ["Date", "Plan", "Calories (kcal)", "Protein (g)", "Carbs (g)", "Fat (g)"],
        "libre": "Free day",
        "principal": "Main plan",
        "dias": ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"],
        "orden": "The start date can't be after the end date.",
        "largo": f"The range can't be longer than {MAX_DIAS} days.",
    },
}


def _textos(idioma: str) -> Dict:
    return TEXTOS.get(idioma, TEXTOS["es"])


def validar_rango(desde: date, hasta: date, idioma: str = "es") -> None:
    tx = _textos(idioma)
    if desde > hasta:
        raise ValueError(tx["orden"])
    if (hasta - desde).days + 1 > MAX_DIAS:
        raise ValueError(tx["largo"])


def _relleno(consumido: float, meta: float):
    if not consumido or not meta:
        return None
    desvio = abs(consumido / meta - 1)
    return FILL_OK if desvio <= 0.10 else FILL_AVISO if desvio <= 0.20 else FILL_MAL


def resumen_por_dia(usuario_id: int, desde: date, hasta: date) -> List[Dict]:
    """Totales consumidos y metas del plan de cada día entre `desde` y `hasta` (ambos incluidos)."""
    fechas = [(desde + timedelta(days=i)).isoformat() for i in range((hasta - desde).days + 1)]
    db = SessionLocal()
    try:
        registros = db.query(RegistroConsumoDB).filter(
            RegistroConsumoDB.usuario_id == usuario_id,
            RegistroConsumoDB.fecha >= fechas[0],
            RegistroConsumoDB.fecha <= fechas[-1],
        ).all()
        planes = plan_service.planes_por_fecha(db, usuario_id, fechas)
        libres = dia_libre_service.fechas_libres(db, usuario_id, fechas)
    finally:
        db.close()

    totales = {f: dict.fromkeys(MACROS, 0.0) for f in fechas}
    for r in registros:
        if r.fecha in totales:
            for m in MACROS:
                totales[r.fecha][m] += getattr(r, m) or 0
    dias = []
    for f in fechas:
        plan = planes[f]
        dias.append({
            "fecha": f,
            "plan": plan.nombre,
            "libre": f in libres,
            "consumido": totales[f],
            "meta": {m: getattr(plan, f"{m}_objetivo") or 0 for m in MACROS},
        })
    return dias


def generar_reporte(usuario_id: int, desde: date, hasta: date, idioma: str = "es") -> bytes:
    """Genera el .xlsx del rango de fechas para un usuario y devuelve sus bytes."""
    validar_rango(desde, hasta, idioma)
    tx = _textos(idioma)
    dias = resumen_por_dia(usuario_id, desde, hasta)

    wb = Workbook()
    ws = wb.active
    ws.title = tx["hoja"]
    ws["A1"] = tx["titulo"]
    ws["A1"].font = Font(bold=True, size=16, color=COLOR_TITULO)
    ws["A2"] = tx["rango"].format(desde=desde.strftime("%d/%m/%Y"), hasta=hasta.strftime("%d/%m/%Y"))
    ws["A2"].font = Font(italic=True, size=10, color="5E5E5E")

    fila_cab = 4
    for col, titulo in enumerate(tx["cabecera"], start=1):
        celda = ws.cell(row=fila_cab, column=col, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", start_color=COLOR_CABECERA, end_color=COLOR_CABECERA)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    fila = fila_cab
    for dia in dias:
        fila += 1
        f = date.fromisoformat(dia["fecha"])
        ws.cell(row=fila, column=1, value=f"{tx['dias'][f.weekday()]} {f.strftime('%d/%m/%Y')}")
        # El nombre por defecto del plan principal se traduce; los que pone el usuario, no
        plan = tx["principal"] if dia["plan"] == plan_service.NOMBRE_PLAN_PRINCIPAL else dia["plan"]
        ws.cell(row=fila, column=2, value=f"{plan} · {tx['libre']}" if dia["libre"] else plan)
        for col, m in enumerate(MACROS, start=3):
            consumido, meta = round(dia["consumido"][m]), round(dia["meta"][m])
            celda = ws.cell(row=fila, column=col, value=f"{consumido}/{meta}")
            celda.alignment = Alignment(horizontal="center")
            relleno = FILL_LIBRE if dia["libre"] else _relleno(consumido, meta)
            if relleno:
                celda.fill = relleno

    tabla = Table(displayName="TablaResumen", ref=f"A{fila_cab}:F{fila}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False)
    ws.add_table(tabla)
    ws.freeze_panes = f"A{fila_cab + 1}"
    for i, ancho in enumerate([16, 26, 17, 15, 19, 13], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
